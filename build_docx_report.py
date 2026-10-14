# -*- coding: utf-8 -*-
"""
build_docx_report.py
Generator Dokumen Word Resmi: Laporan_Desain_Dimensional_NRG.docx
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
DOCX_FILE = os.path.join(OUTPUT_DIR, "Laporan_Desain_Dimensional_NRG.docx")

# Data Tables
from create_reports import DATA_SUMBER, DATA_SCHEMA, DATA_STTM, DATA_Q1, DATA_Q2_SEGMENT, DATA_Q2_SAMPLES, DATA_Q3, DATA_Q4

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_heading(doc, text, level):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run.font.name = "Calibri"
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79) # Deep Navy
    elif level == 2:
        p.paragraph_format.space_before = Pt(13)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run.font.name = "Calibri"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6) # Steel Blue
    elif level == 3:
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0x13, 0x4F, 0x5C) # Deep Teal
    return p

def add_p(doc, text="", bold_prefix="", italic_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(10)
        r_b.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    if italic_prefix:
        r_i = p.add_run(italic_prefix)
        r_i.italic = True
        r_i.font.name = "Calibri"
        r_i.font.size = Pt(10)
        r_i.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    if text:
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p

def add_bullet(doc, text="", bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(10)
        r_b.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    if text:
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p

def add_callout(doc, text, title="CATATAN TEKNIS"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="1F4E79"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}: ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9)
    r_text.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(2)

def add_code(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="16" w:space="0" w:color="888888"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(code_text.strip())
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(2)

def make_table(doc, col_widths, headers, data, align_rights=[]):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1F4E79")
        set_cell_margins(hdr_cells[i], top=90, bottom=90, left=80, right=80)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if i in align_rights:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, row_data in enumerate(data):
        row_cells = tbl.add_row().cells
        bg_col = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=80, right=80)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if c_idx in align_rights:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    set_table_borders(tbl)
    for row in tbl.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(3)

def generate_docx():
    print("Membangun Dokumen Word: Laporan_Desain_Dimensional_NRG.docx...")
    doc = docx.Document()

    # Set Margins (Normal: 1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ----------------------------------------------------
    # COVER / HEADER
    # ----------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("LAPORAN PENGEMBANGAN DATA WAREHOUSE & BUSINESS INTELLIGENCE")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run("Konsolidasi Data Penjualan Omnichannel Nusantara Retail Group (NRG)")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Metadata Box
    tbl_meta = doc.add_table(rows=4, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Mata Kuliah", "Data Warehouse & Business Intelligence (DWBI)"),
        ("Bentuk Penugasan", "Tugas Kelompok 1 (4 Anggota)"),
        ("Tools & Platform", "Pentaho Data Integration (PDI / Kettle v10.2), MySQL 8.0, Python 3.13"),
        ("Tanggal Selesai", "14 Oktober 2026")
    ]
    for idx, (label, val) in enumerate(meta_info):
        r_cells = tbl_meta.rows[idx].cells
        r_cells[0].text = label
        r_cells[1].text = val
        set_cell_background(r_cells[0], "F0F4F8")
        set_cell_background(r_cells[1], "FFFFFF")
        set_cell_margins(r_cells[0], top=60, bottom=60, left=100, right=100)
        set_cell_margins(r_cells[1], top=60, bottom=60, left=100, right=100)
        r_cells[0].paragraphs[0].runs[0].font.bold = True
        r_cells[0].paragraphs[0].runs[0].font.size = Pt(9.5)
        r_cells[1].paragraphs[0].runs[0].font.size = Pt(9.5)
    set_table_borders(tbl_meta)
    for row in tbl_meta.rows:
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.3)

    doc.add_page_break()

    # ----------------------------------------------------
    # RINGKASAN EKSEKUTIF
    # ----------------------------------------------------
    add_heading(doc, "RINGKASAN EKSEKUTIF (EXECUTIVE SUMMARY)", level=1)
    add_p(doc, "Nusantara Retail Group (NRG) menghadapi tantangan operasional analitik yang signifikan akibat fragmentasi sistem data (data silo) di tiga kanal penjualan: 42 toko fisik, situs e-commerce web resmi, dan dua platform marketplace pihak ketiga. Tim data analitik membutuhkan waktu hingga 3 hari kerja manual di spreadsheet Excel untuk mengonsolidasi laporan performa kuartalan bagi jajaran direksi.")
    add_p(doc, "Laporan ini menyajikan arsitektur dan implementasi Data Warehouse berbasis Skema Bintang (Star Schema) dengan tabel fakta penjualan tunggal (fact_penjualan) berbutir halus (atomic grain: satu baris per item transaksi) yang dikelilingi oleh lima tabel dimensi terstandardisasi (dim_waktu, dim_produk, dim_toko_kanal, dim_pelanggan, dim_promosi) serta tabel audit kualitas data (error_kampanye_marketing). Seluruh proses ETL diimplementasikan secara otomatis menggunakan Pentaho Data Integration (PDI / Kettle) dan Python.")
    add_callout(doc, "Seluruh 14.723 baris fakta penjualan dari kelima sistem sumber (S1-S5) berhasil dimuat dengan rekonsiliasi finansial 100% konsisten (net_sales_mismatch = 0, margin_mismatch = 0). Empat kueri analitik strategis berhasil dieksekusi dengan performa sub-detik untuk menjawab seluruh pertanyaan Direktur Penjualan.", title="PENCAPAIAN UTAMA PROYEK")

    # ----------------------------------------------------
    # BAB 1
    # ----------------------------------------------------
    add_heading(doc, "BAB 1: LATAR BELAKANG & RUMUSAN MASALAH BISNIS", level=1)
    add_heading(doc, "1.1 Profil Nusantara Retail Group (NRG)", level=2)
    add_p(doc, "Nusantara Retail Group (NRG) adalah perusahaan ritel multi-kategori nasional di Indonesia yang mendistribusikan produk Elektronik, Furnitur, dan Pakaian. Dalam melayani pasar, NRG mengoperasikan tiga pilar kanal distribusi:")
    add_bullet(doc, "Terdiri dari 42 gerai modern (TS01 s.d. TS42) yang tersebar di kota-kota besar Indonesia.", bold_prefix="1. Toko Fisik (Offline Retail): ")
    add_bullet(doc, "Toko online mandiri resmi milik korporasi berbasis MySQL OLTP.", bold_prefix="2. Situs Web Resmi (Web D2C): ")
    add_bullet(doc, "Akun toko resmi di dua marketplace eksternal terkemuka (Marketplace A dan B).", bold_prefix="3. Marketplace Pihak Ketiga: ")

    add_heading(doc, "1.2 Fenomena Silo Data & Beban Operasional Pelaporan", level=2)
    add_p(doc, "Masing-masing kanal beroperasi pada infrastruktur teknologi yang berbeda dan terisolasi: POS toko fisik berbasis SQL Server diekspor ke CSV titik-koma, e-commerce menggunakan MySQL, marketplace diakses melalui REST API JSON berpaginasi, master produk berada di ERP pusat, dan promosi dicatat manual di spreadsheet marketing.")
    add_p(doc, "Ketiadaan integrasi data terpusat menyebabkan laporan bisnis kuartalan harus dikerjakan secara manual di Excel selama sekitar 3 hari kerja, membebani tim analitik dan memperlambat pengambilan keputusan strategis.")

    add_heading(doc, "1.3 Empat Pertanyaan Kunci Direktur Penjualan", level=2)
    add_p(doc, "Direktur Penjualan NRG mengajukan empat pertanyaan strategis yang harus dapat dijawab secara instan oleh data warehouse:")
    add_bullet(doc, "Berapa margin kotor per sub-kategori produk per kanal dibandingkan dengan kuartal yang sama tahun lalu (YoY)?", bold_prefix="Pertanyaan 1: ")
    add_bullet(doc, "Apakah pelanggan online di situs web adalah orang yang sama dengan pemegang kartu member di toko fisik?", bold_prefix="Pertanyaan 2: ")
    add_bullet(doc, "Kampanye diskon mana yang berhasil menaikkan volume penjualan tanpa menggerus margin keuntungan?", bold_prefix="Pertanyaan 3: ")
    add_bullet(doc, "Produk apa saja yang penjualannya tinggi di marketplace tetapi rendah di toko fisik dan web, sehingga perlu dialihkan stoknya?", bold_prefix="Pertanyaan 4: ")

    # ----------------------------------------------------
    # BAB 2
    # ----------------------------------------------------
    add_heading(doc, "BAB 2: TINJAUAN SISTEM SUMBER DATA & MITIGASI KUALITAS DATA", level=1)
    add_p(doc, "Berikut ringkasan kelima sistem sumber data beserta bentuk berkas dan tantangan kualitas data yang ditangani di pipeline:")
    make_table(doc, [0.6, 1.2, 1.3, 1.6, 1.8], ["Kode", "Sistem", "Kanal / Peran", "Bentuk Berkas", "Keterangan / Tantangan Kualitas"], DATA_SUMBER)

    add_heading(doc, "2.1 Penanganan Anomali Kualitas Data Kritis", level=2)
    add_bullet(doc, "Harga unit_price pada S1 dan S3 sudah termasuk PPN 11%, sedangkan S2 belum termasuk PPN. Pipeline menstandarkan seluruh harga menjadi non-PPN dengan formula unit_price / 1.11.", bold_prefix="Penyeragaman Basis Pajak: ")
    add_bullet(doc, "Waktu transaksi paid_at pada S3 berformat UTC (Z). Pipeline menambahkan +7 jam untuk menyeragamkannya ke Waktu Indonesia Barat (WIB).", bold_prefix="Penyeragaman Basis Waktu: ")
    add_bullet(doc, "Transaksi kasir berstatus 'VOID' pada S2 dibersihkan dan dibuang dari pipeline fakta.", bold_prefix="Pembersihan Transaksi Batal: ")
    add_bullet(doc, "Transaksi ORD-9999 di S1 memuat produk SKU-NEWARR-001 yang belum tercatat di ERP S4. Ditangani melalui Inferred Member P-1061 pada dim_produk.", bold_prefix="Late-Arriving Dimension: ")
    add_bullet(doc, "Nilai target_diskon_persen berupa teks 'sepuluh' pada S5 dialirkan ke error_kampanye_marketing untuk audit tim marketing.", bold_prefix="Penanganan Data Teks 'sepuluh': ")
    add_bullet(doc, "Sekitar 2% transaksi retry pemanggilan API pada S3 dideduplikasi berdasarkan invoice + seller_sku teranyar.", bold_prefix="Deduplikasi API Marketplace: ")

    # ----------------------------------------------------
    # BAB 3
    # ----------------------------------------------------
    add_heading(doc, "BAB 3: DESAIN DIMENSIONAL (DIMENSIONAL DESIGN)", level=1)
    add_heading(doc, "3.1 Struktur Skema Bintang (Star Schema)", level=2)
    add_p(doc, "Model dimensional menggunakan satu tabel fakta penjualan (fact_penjualan) berbutir halus yang dikelilingi oleh lima tabel dimensi terkonformasi:")
    make_table(doc, [1.5, 1.2, 1.5, 0.7, 1.6], ["Nama Tabel", "Jenis", "Primary Key", "Jumlah Baris", "Deskripsi Entitas"], DATA_SCHEMA)

    add_heading(doc, "3.2 Grain dan Formula Perhitungan Fakta", level=2)
    add_bullet(doc, "Satu baris per item transaksi yang dibeli (order line, trx line, invoice item).", bold_prefix="Grain Fakta: ")
    add_bullet(doc, "Seluruh metrik moneter dinyatakan bersih tanpa PPN 11% dalam zona WIB.", bold_prefix="Basis Finansial: ")
    add_bullet(doc, "margin_kotor = nilai_bersih_excl - total_harga_pokok - fee_komisi. Fee komisi hanya berlaku untuk kanal Marketplace dan dialokasikan proporsional per item.", bold_prefix="Formula Margin Kotor: ")

    add_heading(doc, "3.3 Source-to-Target Mapping (STTM)", level=2)
    add_p(doc, "Berikut pemetaan kolom demi kolom dari sistem sumber ke skema bintang target:")
    make_table(doc, [1.1, 1.1, 0.9, 0.8, 1.1, 1.5], ["Target Table", "Target Column", "Tipe Data", "Source", "Source Column", "Logika Transformasi"], DATA_STTM)

    # ----------------------------------------------------
    # BAB 4
    # ----------------------------------------------------
    add_heading(doc, "BAB 4: ARSITEKTUR PIPELINE ETL PENTAHO DATA INTEGRATION", level=1)
    add_p(doc, "Pemuatan data warehouse diorkestrasi oleh master job Pentaho (main_etl_job.kjb) yang menjalankan 8 transformasi (.ktr) secara sekuensial:")
    add_bullet(doc, "Mengenerasi 1.096 hari (2022-2024), menghitung kuartal, bulan, hari, dan surrogate key YYYYMMDD.", bold_prefix="1. dim_waktu_loader.ktr: ")
    add_bullet(doc, "Membaca produk_master.csv (S4), memilih HPP mutakhir, menginjeksi inferred member P-1061 untuk ORD-9999 (61 produk).", bold_prefix="2. dim_produk_loader.ktr: ")
    add_bullet(doc, "Membaca store_master.csv (S2), standarisasi kota via ref_kota.csv, menyuntikkan 3 kanal online (45 entitas).", bold_prefix="3. dim_toko_kanal_loader.ktr: ")
    add_bullet(doc, "Membaca customers (S1) dan member.csv (S2), normalisasi nomor HP, resolusi identitas omnichannel (2.345 profil).", bold_prefix="4. dim_pelanggan_loader.ktr: ")
    add_bullet(doc, "Trimming spasi, filter data kotor, mengalirkan error 'sepuluh' ke tabel audit, memuat 11 promosi valid.", bold_prefix="5. dim_promosi_loader.ktr: ")
    add_bullet(doc, "Normalisasi PPN 11%, kalkulasi margin kotor dan COGS, memuat 4.540 transaksi web ke fact_penjualan.", bold_prefix="6. fact_penjualan_loader.ktr: ")
    add_bullet(doc, "Membaca 42 toko fisik, memfilter transaksi berstatus VOID pada header, memuat 8.386 baris transaksi kasir.", bold_prefix="7. fact_s2_pos_loader.ktr: ")
    add_bullet(doc, "Membaca 16 berkas JSON marketplace, membongkar nested items, deduplikasi unik stempel terbaru, konversi WIB (+7 jam), memuat 1.797 baris.", bold_prefix="8. fact_s3_mpl_loader.ktr: ")
    add_callout(doc, "Seluruh 8 transformasi Pentaho (.ktr) dan master job (main_etl_job.kjb) telah diverifikasi berjalan dengan status EXIT CODE 0 tanpa error di Pentaho Kitchen dan Spoon, memuat tepat 14.723 baris fakta. Tersedia pula script run_etl.py untuk pemuatan otomatis dalam 3,92 detik.", title="PERFORMA & VALIDASI ETL ENGINE")

    # ----------------------------------------------------
    # BAB 5
    # ----------------------------------------------------
    add_heading(doc, "BAB 5: ANALISIS HASIL EKSEKUSI EMPAT KUERI PEMBUKTIAN BISNIS", level=1)

    # 5.1
    add_heading(doc, "5.1 Pertanyaan 1: Analisis Margin YoY Q1 2023 vs Q1 2022 per Sub-Kategori & Kanal", level=2)
    add_p(doc, "Kueri SQL membandingkan margin kotor Q1 2023 terhadap Q1 2022 untuk seluruh sub-kategori produk dan kanal penjualan:")
    make_table(doc, [1.4, 1.0, 0.7, 1.1, 1.1, 0.7], ["Sub-Kategori", "Kanal", "Kuartal", "Margin Q1 2023", "Margin Q1 2022", "YoY %"], [[r[0], r[1], r[2], r[3], r[4], r[6]] for r in DATA_Q1], align_rights=[3, 4, 5])
    add_p(doc, "Analisis Kritis: Sub-kategori Kursi mengalami kenaikan margin tertinggi di Web (+37,03%) dan Marketplace (+25,69%). Lemari di Toko Fisik menyumbang kenaikan laba terbesar (+Rp 168,84 Juta). Sebaliknya, Meja di Marketplace anjlok -54,77% dan Sepatu di Web turun -31,77%.")

    # 5.2
    add_heading(doc, "5.2 Pertanyaan 2: Analisis Pelanggan Omnichannel vs Online Murni", level=2)
    add_p(doc, "Hasil segmentasi pelanggan yang berbelanja di kanal situs web resmi:")
    make_table(doc, [1.8, 0.7, 0.7, 0.7, 1.1, 1.0], ["Segmen Pelanggan", "User", "Trx Web", "Unit", "Total Belanja (IDR)", "Total Margin (IDR)"], [[r[0], r[1], r[2], r[3], r[4], r[6]] for r in DATA_Q2_SEGMENT], align_rights=[1, 2, 3, 4, 5])

    add_p(doc, "Sampel Profil 5 Pelanggan Omnichannel Teratas Berbelanja Lintas Kanal:")
    make_table(doc, [0.6, 1.5, 1.1, 0.6, 0.6, 1.1, 1.0], ["SK", "Nama Pelanggan", "No HP Normal", "Web", "Toko", "Belanja Web", "Total Belanja"], [[r[0], r[1], r[2], r[4], r[5], r[6], r[7]] for r in DATA_Q2_SAMPLES[:5]], align_rights=[3, 4, 5, 6])
    add_p(doc, "Analisis Kritis: Sebanyak 200 pelanggan web teridentifikasi secara pasti sebagai member toko fisik (199 aktif bertransaksi). Segmen omnichannel ini adalah High-Value Customer yang berbelanja puluhan kali lintas kanal dengan kontribusi belanja gabungan ratusan juta rupiah per pelanggan.")

    # 5.3
    add_heading(doc, "5.3 Pertanyaan 3: Evaluasi Efektivitas Kampanye Promosi terhadap Volume & Margin", level=2)
    add_p(doc, "Evaluasi perbandingan volume penjualan dan margin kotor antar promosi:")
    make_table(doc, [1.0, 1.4, 0.6, 0.6, 0.6, 1.1, 0.7, 1.0], ["Kode", "Nama Kampanye", "Diskon", "Trx", "Unit", "Net Sales (IDR)", "Margin %", "Evaluasi"], [[r[0], r[1], r[3], r[4], r[5], r[6], r[9], r[10]] for r in DATA_Q3], align_rights=[3, 4, 5, 6])
    add_p(doc, "Analisis Kritis: Promo FLASH50 (5%) dan BACKSCHOOL (10%) terbukti sangat efektif menjaga margin sehat (19,83% dan 18,52%). Sebaliknya, HARBOLNAS (20%) menggerus margin hingga 13,64% (turun lebih dari 5% dari batas normal), membuktikan bahwa diskon 20% tanpa syarat belanja minimum sangat merugikan profitabilitas.")

    # 5.4
    add_heading(doc, "5.4 Pertanyaan 4: Deteksi Kanibalisasi & Produk Unggulan Marketplace untuk Realokasi Stok", level=2)
    add_p(doc, "Top produk dengan pangsa pasar penjualan tertinggi di marketplace:")
    make_table(doc, [0.7, 1.8, 0.8, 0.6, 0.6, 0.6, 0.7, 0.7], ["Kode", "Nama Produk", "Kategori", "MPL", "Toko", "Web", "Pangsa MPL", "Margin MPL"], [[r[0], r[1], r[2], r[4], r[5], r[6], r[8], r[10]] for r in DATA_Q4], align_rights=[3, 4, 5, 6, 7])
    add_p(doc, "Analisis Kritis: Produk P-1019 (Samsung Smartphone 128GB) dan P-1054 (Uniqlo Atasan) mendominasi penjualan marketplace dengan pangsa pasar >20% secara nasional. Keduanya adalah kandidat utama realokasi pasokan stok dari toko fisik ke gudang fulfillment marketplace.")

    # ----------------------------------------------------
    # BAB 6
    # ----------------------------------------------------
    add_heading(doc, "BAB 6: KESIMPULAN & REKOMENDASI STRATEGIS BISNIS", level=1)
    add_bullet(doc, "Data warehouse berbasis Star Schema berhasil menyatukan 5 sistem sumber heterogen menjadi Single Source of Truth yang andal dan konsisten.", bold_prefix="1. Kesimpulan Teknis: ")
    add_bullet(doc, "Batasi diskon agresif 20% tanpa batas minimal belanja. Perbanyak promo diskon terukur 5%-10% seperti FLASH50 dan BACKSCHOOL.", bold_prefix="2. Kebijakan Diskon: ")
    add_bullet(doc, "Integrasikan program loyalitas member di mana poin belanja toko fisik dapat dipakai di web dan sebaliknya untuk mempertahankan segmen High-Value Omnichannel.", bold_prefix="3. Loyalitas Omnichannel: ")
    add_bullet(doc, "Pindahkan 15%-20% alokasi stok produk P-1019 dan P-1054 ke fulfillment center marketplace guna mempercepat pengiriman dan menghindari stock-out.", bold_prefix="4. Realokasi Rantai Pasok: ")

    # ----------------------------------------------------
    # BAB 7
    # ----------------------------------------------------
    add_heading(doc, "BAB 7: PANDUAN INSTALASI & EKSEKUSI PIPELINE DARI KONDISI KOSONG", level=1)
    add_p(doc, "Langkah-langkah menjalankan data warehouse dari kondisi kosong:")
    add_code(doc, """# 1. Inisialisasi Database MySQL
mysql -u root -e "CREATE DATABASE IF NOT EXISTS nrg_dw;"
mysql -u root nrg_dw < dw_schema_ddl.sql
mysql -u root -e "CREATE DATABASE IF NOT EXISTS nrg_ecom;"
mysql -u root nrg_ecom < Data\\S1_ecom_db_schema_mysql.sql
cd Data && python load_all.py && cd ..

# 2. Jalankan Master ETL Pentaho
kitchen.bat /file:"c:\\Users\\daffa\\Downloads\\Tugas 1 DWBI KELOMPOK\\main_etl_job.kjb" /level:Basic

# 3. Atau Jalankan Python ETL Runner Cepat (3.92 detik)
python run_etl.py

# 4. Jalankan Kueri Pembuktian
mysql -u root nrg_dw < kueri_pembuktian.sql""")

    doc.save(DOCX_FILE)
    print(f"Sukses! Dokumen Word tersimpan di: {DOCX_FILE}")

if __name__ == "__main__":
    generate_docx()
